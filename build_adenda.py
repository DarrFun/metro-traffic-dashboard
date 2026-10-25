# -*- coding: utf-8 -*-
"""Genera la adenda IEEE (Sección V-B) en formato .docx."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path

BASE = Path(__file__).resolve().parent
FIGS = BASE / "figs"
OUT = BASE / "docs" / "adenda_ieee_seccion_VB.docx"
OUT.parent.mkdir(exist_ok=True)

doc = Document()

# ---- estilos base ----
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

def set_heading_style(style_name, size, bold=True, color=(0, 0, 0), space_before=10, space_after=4):
    st = doc.styles[style_name]
    st.font.name = "Times New Roman"
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = RGBColor(*color)
    st.paragraph_format.space_before = Pt(space_before)
    st.paragraph_format.space_after = Pt(space_after)
    st.paragraph_format.keep_with_next = True

set_heading_style("Heading 1", 13)
set_heading_style("Heading 2", 11.5)
set_heading_style("Heading 3", 10.5)

def para(text, bold=False, italic=False, size=10.5, align=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    return p

def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    return p

def figure(path, caption, width=6.3):
    if not path.exists():
        para(f"[Figura no encontrada: {path.name}]", italic=True, size=9)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    cap.paragraph_format.space_after = Pt(10)

# ---- título de la adenda ----
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Adenda al Artículo IEEE — Sección V-B")
r.bold = True
r.font.size = Pt(14)
title.paragraph_format.space_after = Pt(2)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Arquitectura Visual y Diseño de la Interfaz Analítica")
r.bold = True
r.font.size = Pt(11.5)
sub.paragraph_format.space_after = Pt(2)

auth = doc.add_paragraph()
auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = auth.add_run("Jonny Antonio Quintanilla Yanes, Darren Steve Funes Rosales, Josue Adony Morales Torres, "
                 "Irvin Alejandro Arguera Canjura y Rodrigo Ernesto Grande Zelaya")
r.font.size = Pt(9.5)
r.italic = True
auth.paragraph_format.space_after = Pt(4)

aff = doc.add_paragraph()
aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = aff.add_run("Facultad de Ingeniería, Universidad Evangélica de El Salvador (UEES) — Asignatura: Análisis de Datos")
r.font.size = Pt(9)
aff.paragraph_format.space_after = Pt(10)

para("Resumen— Esta adenda extiende la Sección V del artículo principal documentando la "
     "traducción del pipeline de datos (Fases I–V) en una interfaz analítica interactiva de baja "
     "carga cognitiva. Se justifican teóricamente los canales visuales elegidos, se presenta la "
     "taxonomía de tareas visuales soportadas, las decisiones de diseño basadas en Gestalt y Tufte, "
     "el flujo de interactividad según la heurística de Shneiderman y la narrativa de datos "
     "(Data Storytelling) que articula el dashboard. Se incluyen figuras anotadas de casos de uso reales.",
     italic=True, size=10)

# ---- V-B ----
doc.add_heading("V-B. Arquitectura Visual y Diseño de la Interfaz", level=1)

doc.add_heading("V-B.1 Justificación teórica de canales visuales", level=2)
para("La interfaz se diseñó siguiendo la Gramática de los Gráficos [9], que separa la "
     "representación en marcas (puntos, líneas, barras, áreas) y canales (posición, longitud, "
     "ángulo, área, color, saturación, forma). La elección de canales se rigió por la "
     "efectividad perceptiva demostrada experimentalmente [10]:")
bullet("Posición y longitud son los canales de mayor precisión para comparaciones cuantitativas; "
       "por ello las comparaciones de tráfico por hora, día, mes y condición meteorológica se "
       "codifican con barras y líneas (longitud/posición vertical).")
bullet("El color y la saturación se usan únicamente con propósito funcional o de alerta: azul "
       "para la variable objetivo (tráfico), naranja para temperatura, rojo para anomalías/feriados "
       "y gris para referencias neutras. No se emplea color decorativo.")
bullet("Se prohíben los elementos que distorsionan la percepción: no se usan gráficos 3D, "
       "gráficos circulares densos ni ejes truncados sin justificación analítica. Las escalas de "
       "los ejes parten de cero cuando se comparan magnitudes.")
bullet("Para la relación entre dos variables continuas (temperatura vs. tráfico) se usa un "
       "diagrama de dispersión con línea de tendencia (OLS), y para la asociación entre variables "
       "se emplea una matriz de correlación de Spearman (canal de color con escala centrada en 0).")

doc.add_heading("V-B.2 Taxonomía de tareas visuales", level=2)
para("El dashboard soporta las tareas visuales identificadas por Munzner [11] para datos "
     "tabulares con dimensión temporal:")
bullet("Búsqueda y recuperación de valores (lookup): KPIs de cabecera (registros, tráfico medio, "
       "máximo, hora pico, temperatura media, feriados).")
bullet("Comparación: tráfico por hora, por día de la semana, por mes y por estación (barras/líneas).")
bullet("Tendencia: serie horaria del tráfico medio con anotación de los picos de la mañana y la tarde.")
bullet("Distribución: histograma del volumen de tráfico (bimodal) y boxplots por condición meteorológica.")
bullet("Correlación: matriz de Spearman y dispersión temperatura–tráfico.")
bullet("Detección de anomalías: contraste feriados vs. días normales y auditoría pre/post (outliers).")

doc.add_heading("V-B.3 Reducción de carga cognitiva (Gestalt y Tufte)", level=2)
para("Se aplicaron los principios de Gestalt y el concepto de data-to-ink ratio de Tufte [12] "
     "para minimizar la carga cognitiva:")
bullet("Maximización del ratio señal/ruido: eliminación de chartjunk, bordes innecesarios, "
       "cuadrículas redundantes, leyendas duplicadas y saturación tipográfica.")
bullet("Ley de proximidad: los KPIs se agrupan en una fila superior; los filtros se concentran en "
       "la barra lateral; las vistas se organizan por pestañas temáticas (Diagnóstico, Patrones "
       "temporales, Impacto meteorológico, Explorador).")
bullet("Ley de similitud: color y forma consistentes para la misma variable en todas las vistas "
       "(azul = tráfico, naranja = temperatura).")
bullet("Ley de cierre: cada gráfico es una unidad visual autocontenida con título, ejes etiquetados "
       "y una línea de caption interpretativa.")
bullet("Jerarquía tipográfica clara: títulos en negrita, subtítulos por sección y captions en "
       "itálica de menor tamaño para la interpretación.")

doc.add_heading("V-B.4 Flujos de interactividad (heurística de Shneiderman)", level=2)
para("La interactividad sigue la tríada overview first → zoom and filter → details-on-demand [13]:")
bullet("Overview first: tarjetas de KPIs y panorámica general (distribución del tráfico, auditoría "
       "pre/post) en la primera pestaña.")
bullet("Zoom and filter: filtros cruzados globales y sincronizados en la barra lateral (rango de "
       "fechas, rango de hora, día de la semana, condición meteorológica, estación, feriados, "
       "horas pico). Todos los gráficos reaccionan al mismo subconjunto.")
bullet("Details-on-demand: tooltips enriquecidos en los gráficos (fecha, tráfico, temperatura, clima), "
       "tabla de inspección de los primeros 200 registros y descarga del subconjunto filtrado en CSV.")
para("El estado filtrado se muestra siempre (registros filtrados y cobertura porcentual), de modo "
     "que el usuario conoce el soporte estadístico de cada vista.")

doc.add_heading("V-B.5 Narrativa de datos (Data Storytelling)", level=2)
para("La interfaz articula un flujo guiado que transforma datos en decisión:")
bullet("Diagnóstico inicial: calidad de los datos (sin nulos ni outliers imposibles tras el pipeline) "
       "y auditoría pre/post de skewness y curtosis.")
bullet("Detección de patrones y anomalías: patrón bimodal por hora (picos ~07:00 y ~16:00), "
       "reducción del tráfico los fines de semana y en feriados.")
bullet("Análisis de causa raíz: el factor temporal (hora) es predictor más fuerte que el meteorológico "
       "(ρ de Spearman 0.33 vs. 0.14 para la temperatura); el clima modula pero no determina el tráfico.")
bullet("Recomendación de acción/decisión: priorizar variables temporales en modelos predictivos, "
       "usar escalado robusto (RobustScaler) para precipitación y considerar transformaciones "
       "logarítmicas para las variables cero-infladas.")

# ---- Figuras anotadas ----
doc.add_page_break()
doc.add_heading("Figuras anotadas (casos de uso reales)", level=2)

figure(FIGS / "fig01_auditoria_prepost.png",
       "Fig. 4. Auditoría pre/post: distribuciones de temp y rain_1h antes (crudo, naranja) y "
       "después (limpio, azul) de la corrección de outliers. El ruido extremo deja de dominar "
       "la distribución.", width=6.4)

figure(FIGS / "fig02_trafico_por_hora.png",
       "Fig. 5. Volumen de tráfico medio por hora del día (canal posición/longitud). Se anotan "
       "los dos picos de congestión (~07:00 y ~16:00), evidencia del patrón bimodal.", width=6.4)

figure(FIGS / "fig03_heatmap_hora_dia.png",
       "Fig. 6. Mapa de calor hora × día de la semana (canal de color, escala secuencial). "
       "Zonas oscuras = picos de congestión en días laborales.", width=6.4)

figure(FIGS / "fig04_boxplot_clima.png",
       "Fig. 7. Distribución del tráfico por condición meteorológica (boxplot, canal posición). "
       "Permite comparar medianas y dispersión entre categorías.", width=6.4)

figure(FIGS / "fig05_scatter_temp.png",
       "Fig. 8. Dispersión temperatura (°C) vs. volumen de tráfico con línea de tendencia OLS "
       "(canal posición + color semántico por temperatura).", width=6.4)

figure(FIGS / "fig06_correlacion.png",
       "Fig. 9. Matriz de correlación de Spearman del dataset procesado (canal de color centrado "
       "en 0). Confirma que la hora es el predictor más fuerte.", width=5.6)

figure(FIGS / "fig07_feriado.png",
       "Fig. 10. Tráfico medio en días normales vs. feriados (canal longitud + color funcional: "
       "azul = normal, rojo = alerta).", width=5.0)

# ---- conclusiones ----
doc.add_heading("Conclusiones de la interfaz", level=2)
para("La interfaz resultante traduce el pipeline de datos en una herramienta analítica accionable: "
     "maximiza el ratio señal/ruido, usa canales visuales de alta efectividad perceptiva, ofrece "
     "interactividad fluida (overview → filter → details) y articula una narrativa que va del "
     "diagnóstico a la recomendación. El dashboard es reproducible (repositorio con requirements.txt "
     "y datos limpios) y se despliega con Streamlit, cumpliendo el requisito de enlace interactivo.")
para("Enlace interactivo (tras despliegue): [URL a completar tras el despliegue en Streamlit Community Cloud].",
     italic=True, size=9.5)

# ---- referencias adicionales ----
doc.add_heading("Referencias (adenda)", level=2)
refs = [
    "[9] L. Wilkinson, The Grammar of Graphics, 2nd ed. New York, NY, USA: Springer, 2005.",
    "[10] W. S. Cleveland and R. McGill, “Graphical perception: Theory, experimentation, and "
    "application to the development of graphical methods,” J. Amer. Statist. Assoc., vol. 79, "
    "no. 387, pp. 531–554, 1984.",
    "[11] T. Munzner, Visualization Analysis and Design. Boca Raton, FL, USA: CRC Press, 2014.",
    "[12] E. R. Tufte, The Visual Display of Quantitative Information, 2nd ed. Cheshire, CT, USA: "
    "Graphics Press, 2001.",
    "[13] B. Shneiderman, “The eyes have it: A task by data type taxonomy for information "
    "visualizations,” in Proc. IEEE Symp. Visual Languages, 1996, pp. 336–343.",
    "[14] C. Ware, Information Visualization: Perception for Design, 3rd ed. Waltham, MA, USA: "
    "Morgan Kaufmann, 2012.",
]
for rtxt in refs:
    p = doc.add_paragraph(rtxt)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(2)
    for run in p.runs:
        run.font.size = Pt(9)

doc.save(OUT)
print("saved", OUT)
